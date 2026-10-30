"""Builds the supplementary report (.docx) from the outputs of the three experiments.

Contents (only these analyses): 1 four-configuration ablation, 2 threshold sensitivity,
3 score distributions. Tables and the interpretation sentences are generated from the numbers in
results/supplementary/. Figures: the MATLAB versions are used when present in
results/supplementary/figures_matlab/ (same file names), otherwise the Python versions.

Usage (Windows, project folder):
  .venv\\Scripts\\pip.exe install python-docx          (once)
  .venv\\Scripts\\python.exe supplementary\\make_report.py
"""
from __future__ import annotations

import argparse
import json
import math
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import common as C  # noqa: E402

NAVY = "1F3864"
CONFIG_LABELS = {"no_defense": "No defense", "detector_only": "Detector only", "sanitizer_only": "Sanitizer only",
                 "hybrid": "Hybrid"}
SPLIT_LABELS = {"regular": "Regular test", "challenge": "Challenge", "all": "Both"}


def p(x, nd=1):
    return "n/a" if x is None or x == "" or (isinstance(x, float) and math.isnan(x)) else f"{100 * float(x):.{nd}f}%"


def pv(x, nd=1):
    """Percentage without the % sign (for table cells whose header states the unit)."""
    return "n/a" if x is None or x == "" or (isinstance(x, float) and math.isnan(x)) else f"{100 * float(x):.{nd}f}"


def ci(lo, hi, nd=1):
    return "" if lo in (None, "") or (isinstance(lo, float) and math.isnan(lo)) else f"[{100 * float(lo):.{nd}f}, {100 * float(hi):.{nd}f}]"


class Report:
    def __init__(self, title: str):
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm, Pt, RGBColor

        self.Pt, self.Cm, self.RGB, self.ALIGN = Pt, Cm, RGBColor, WD_ALIGN_PARAGRAPH
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Cm(2.3))
        normal = self.doc.styles["Normal"]
        normal.font.name, normal.font.size = "Times New Roman", Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        for level, size in ((1, 15), (2, 12.5)):
            st = self.doc.styles[f"Heading {level}"]
            st.font.name, st.font.size, st.font.bold = "Times New Roman", Pt(size), True
            st.font.color.rgb = RGBColor.from_string(NAVY)
        self.figures = self.tables = 0
        t = self.doc.add_paragraph()
        r = t.add_run(title)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(18), RGBColor.from_string(NAVY)

    def h1(self, text):
        self.doc.add_heading(text, level=1)

    def h2(self, text):
        self.doc.add_heading(text, level=2)

    def para(self, text, italic=False, size=None):
        par = self.doc.add_paragraph()
        r = par.add_run(text)
        r.italic = italic
        if size:
            r.font.size = self.Pt(size)
        par.paragraph_format.alignment = self.ALIGN.JUSTIFY
        return par

    def bullets(self, items):
        for text in items:
            self.doc.add_paragraph(text, style="List Bullet")

    def figure(self, path: Path, caption: str, width_cm=11.5):
        self.figures += 1
        par = self.doc.add_paragraph()
        par.alignment = self.ALIGN.CENTER
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(str(path), width=self.Cm(width_cm))
        cap = self.doc.add_paragraph()
        cap.alignment = self.ALIGN.CENTER
        r = cap.add_run(f"Figure {self.figures}: ")
        r.bold, r.font.size = True, self.Pt(9.5)
        cap.add_run(caption).font.size = self.Pt(9.5)

    def table(self, title: str, header: list[str], rows: list[list[str]], widths_cm: list[float] | None = None, note=None):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        self.tables += 1
        cap = self.doc.add_paragraph()
        cap.alignment = self.ALIGN.CENTER
        r = cap.add_run(f"Table {self.tables}: ")
        r.bold, r.font.size = True, self.Pt(9.5)
        cap.add_run(title).font.size = self.Pt(9.5)
        cap.paragraph_format.keep_with_next = True
        tab = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        tab.style = "Table Grid"
        tab.autofit = False
        if widths_cm:
            layout = OxmlElement("w:tblLayout")
            layout.set(qn("w:type"), "fixed")
            tab._tbl.tblPr.append(layout)
            for j, w in enumerate(widths_cm):
                tab.columns[j].width = self.Cm(w)
        for i, cells in enumerate([header, *rows]):
            for j, value in enumerate(cells):
                cell = tab.rows[i].cells[j]
                cell.text = ""
                run = cell.paragraphs[0].add_run(str(value))
                run.font.size = self.Pt(9)
                run.bold = i == 0
                if i == 0:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), "E7ECF4")
                    cell._tc.get_or_add_tcPr().append(shd)
                if widths_cm:
                    cell.width = self.Cm(widths_cm[j])
                cell.paragraphs[0].paragraph_format.keep_with_next = i < len(rows)   # keep the table on one page
            split = OxmlElement("w:cantSplit")
            split.set(qn("w:val"), "true")
            tab.rows[i]._tr.get_or_add_trPr().append(split)
        if note:
            n = self.doc.add_paragraph()
            n.add_run(note).font.size = self.Pt(8.5)
        else:
            self.doc.add_paragraph()

    def save(self, path: Path):
        self.doc.save(str(path))


def pick_figure(out: Path, name: str) -> tuple[Path | None, str]:
    for folder, source in ((out / "figures_matlab", "MATLAB"), (out / "figures_python", "Python")):
        if (folder / f"{name}.png").exists():
            return folder / f"{name}.png", source
    return None, ""


# --------------------------------------------------------------------------- section 1: ablation
def ablation_section(rep: Report, out: Path):
    rep.h1("1  Four-configuration ablation")
    f = out / "ablation_summary.json"
    if not f.exists():
        rep.para("Not run yet. Run: .venv\\Scripts\\python.exe supplementary\\ablation_four_configs.py", italic=True)
        return
    meta = json.loads(f.read_text())
    summ = meta["summary"]
    splits = [s for s in meta["splits"]]
    items = meta["items"]
    rep.para("The four configurations answer exactly the same questions over exactly the same chunks with the same "
             "prompt and generator, so differences come from the defense alone.")
    rep.bullets([
        "Items: " + "; ".join(f"{SPLIT_LABELS[s]}: {items[s]['injected']} injected chunks (stratified by attack family) "
                              f"and {items[s]['clean']} clean chunks (the clean twins of injected chunks)" for s in splits) + ".",
        "Context: the chunk the question is about is the only context document (retrieval exposure is guaranteed).",
        "Configurations: no defense (chunk unchanged); detector only (flagged chunks are dropped); sanitizer only "
        "(every chunk is sanitized); hybrid (flagged chunks are sanitized, the others pass unchanged). "
        f"Detector threshold τ = {meta['threshold']:g}.",
        f"Generator: {meta['llm']} via Ollama, temperature {meta['llm_temperature']}, seed {meta['llm_seed']}, "
        f"at most {meta['llm_max_tokens']} tokens; identical prompts are generated once and reused.",
        "Attack success (ASR): the attacker's target phrase appears in the answer. Correct answer: every expected value "
        "of the question (two facts) appears in the answer. Both are lexical proxies; 95% Wilson intervals in brackets.",
        f"Cost of the defense on this CPU: Layer 1 {meta['detector_ms_per_chunk']:.1f} ms per chunk, Layer 2 "
        f"{meta['sanitizer_ms_per_chunk']:.1f} ms per chunk (no extra LLM call in any configuration).",
    ])
    rows = []
    for r in summ:
        if r["split"] not in splits:
            continue
        flagged = (f"{r['flagged_injected']}/{r['n_injected']} | {r['flagged_clean']}/{r['n_clean']}"
                   if r["flagged_injected"] != "" else "—")
        if r["config"] == "detector_only":
            changed = f"dropped {r['dropped_injected']} | {r['dropped_clean']}"
        elif r["changed_injected"] != "":
            changed = f"edited {r['changed_injected']} | {r['changed_clean']}"
        else:
            changed = "—"
        rows.append([SPLIT_LABELS[r["split"]], CONFIG_LABELS[r["config"]],
                     f"{pv(r['asr'])} {ci(r['asr_lo'], r['asr_hi'])}",
                     f"{pv(r['correct_attack_rate'])} {ci(r['correct_attack_lo'], r['correct_attack_hi'])}",
                     f"{pv(r['correct_clean_rate'])} {ci(r['correct_clean_lo'], r['correct_clean_hi'])}",
                     flagged, changed])
    rep.table("Attack success and answer correctness per configuration, in percent with 95% Wilson intervals "
              "(same items in every row of a test set)",
              ["Test set", "Configuration", "ASR (injected)", "Correct (injected)", "Correct (clean)",
               "Flagged by Layer 1 (inj. | clean)", "Chunk changed (inj. | clean)"],
              rows, [1.8, 2.6, 2.4, 2.4, 2.4, 2.4, 2.4])
    fig, source = pick_figure(out, "fig_ablation")
    if fig:
        rep.figure(fig, "Attack success (top) and correct answers on injected (middle) and clean items (bottom) for the "
                        f"four configurations; whiskers are 95% Wilson intervals ({source} figure).", 10.5)
    fam_file = out / "ablation_family_asr.csv"
    if fam_file.exists():
        fam = [r for r in C.read_rows(fam_file) if r["split"] == "challenge"]
        if fam:
            names = sorted({r["family"] for r in fam})
            by = {(r["family"], r["config"]): r for r in fam}
            rows = [[n, str(by.get((n, "no_defense"), {}).get("n", "")),
                     *[p(float(by[(n, c)]["asr"])) if (n, c) in by else "—" for c in CONFIG_LABELS]] for n in names]
            rep.table("Attack success per withheld attack family (challenge set)",
                      ["Family", "n", *CONFIG_LABELS.values()], rows, [4.0, 1.2, 2.8, 2.8, 2.8, 2.8])
    rep.h2("Interpretation")
    for text in interpret_ablation(summ, splits):
        rep.para(text)


def interpret_ablation(summ, splits):
    get = {(r["split"], r["config"]): r for r in summ}
    out = []
    for s in splits:
        if not all((s, c) in get for c in CONFIG_LABELS):
            continue
        g = {c: get[(s, c)] for c in CONFIG_LABELS}
        asr = {c: g[c]["asr"] for c in g}
        att = {c: g[c]["correct_attack_rate"] for c in g}
        cln = {c: g[c]["correct_clean_rate"] for c in g}
        best_asr = min(asr.values())
        lowest = [CONFIG_LABELS[c] for c in asr if abs(asr[c] - best_asr) < 1e-12]
        best_att = max(att.values())
        top_att = [CONFIG_LABELS[c] for c in att if abs(att[c] - best_att) < 1e-12]
        text = (f"{SPLIT_LABELS[s]}: without a defense the generator produced the attacker's phrase for {p(asr['no_defense'])} "
                f"of the injected items and answered {p(att['no_defense'])} of them correctly. Detector only gave ASR "
                f"{p(asr['detector_only'])}, sanitizer only {p(asr['sanitizer_only'])} and the hybrid "
                f"{p(asr['hybrid'])}; the lowest ASR was reached by {' and '.join(lowest)}. Under attack, answers were "
                f"correct for {p(att['detector_only'])} (detector only), {p(att['sanitizer_only'])} (sanitizer only) and "
                f"{p(att['hybrid'])} (hybrid); the most correct answers came from {' and '.join(top_att)}.")
        dropped = g["detector_only"]["dropped_injected"]
        if dropped:
            text += (f" Detector only removes the whole chunk when it is flagged ({dropped} of {g['detector_only']['n_injected']} "
                     "injected chunks), so the policy facts disappear together with the injection; the hybrid keeps them "
                     "and removes only the flagged clauses.")
        diff = max(abs(cln[c] - cln["no_defense"]) for c in cln)
        text += (f" On clean chunks, correctness was {p(cln['no_defense'])} without a defense and {p(cln['hybrid'])} with "
                 f"the hybrid" + (", i.e. the defenses did not change benign answers." if diff < 1e-12 else
                                  f"; the largest change from no defense across configurations was {100 * diff:.1f} points."))
        out.append(text)
    out.append("These are lexical proxies on synthetic data: a quoted or refused target phrase counts as an attack "
               "success, and an answer can contain the expected numbers yet still be wrong in other ways.")
    return out


# --------------------------------------------------------------------------- section 2: thresholds
SUP = str.maketrans("-0123456789", "⁻⁰¹²³⁴⁵⁶⁷⁸⁹")


def tfmt(t: float) -> str:
    """Thresholds as written in the text: 0.1 -> 0.10, 0.125 -> 0.125, 1e-06 -> 0.000001, 0.999999."""
    if 0 < t < 0.001 or 0.999 < t < 1:
        return f"{t:.{-math.floor(math.log10(t if t < 0.5 else 1 - t) + 1e-9)}f}"
    return f"{t:.2f}" if abs(t * 100 - round(t * 100)) < 1e-9 else f"{t:g}"


def sfmt(x: float) -> str:
    """A score with three significant digits of its distance to 0 or to 1 (0.000552, 0.175, 0.999842)."""
    if not 0 < x < 1:
        return f"{x:g}"
    return f"{x:.{max(3, 2 - math.floor(math.log10(min(x, 1 - x)) + 1e-9))}f}"


def axis_end(logodds: float) -> str:
    t = 1 / (1 + math.exp(-logodds))
    m, e = f"{min(t, 1 - t):.1e}".split("e")
    sci = f"{m} × 10{str(int(e)).translate(SUP)}"
    return sci if t < 0.5 else f"1 − {sci}"


def margin(x) -> str:
    return "—" if x is None or (isinstance(x, float) and math.isnan(x)) else f"{x:+.2f}"


def roles(text: str) -> str:
    """Role column: 'requested, current, selected' -> 'current, τ*'."""
    names = {"current": "current", "selected": "τ*", "midpoint": "midpoint"}
    return ", ".join(names[r] for r in (x.strip() for x in text.split(",")) if r in names)


def threshold_section(rep: Report, out: Path):
    rep.h1("2  Threshold sensitivity of the detector")
    f = out / "threshold_summary.json"
    if not f.exists():
        rep.para("Not run yet. Run: .venv\\Scripts\\python.exe supplementary\\threshold_sensitivity.py", italic=True)
        return
    s = json.loads(f.read_text())
    if "optimal_interval" not in s:
        rep.para("threshold_summary.json comes from an older version of threshold_sensitivity.py; run the script again.",
                 italic=True)
        return
    v, cur, req = s["validation"], s["current_threshold"], s["thresholds"]
    if s.get("checkpoint_threshold_differs"):
        rep.para(f"Checkpoint note: {Path(s['detector_dir']).name} stores the threshold {s['checkpoint_threshold']}, not "
                 f"{cur:g}. Unless {cur:g} was selected for this checkpoint, this section cannot justify it. Run "
                 "threshold_sensitivity.py --list-checkpoints and repeat the analysis with the checkpoint that selected "
                 f"{cur:g} (--detector-dir <folder> --rescore-tests).", italic=True)
    rep.h2("Method")
    x_lo, x_hi = s["axis"]
    rule = s["rule"].replace("<= limit", f"≤ {100 * s['fpr_limit']:g}%")
    rep.bullets([
        f"Selection data: the validation set only ({v['chunks']:,} chunks: {v['injected']:,} injected, {v['benign']:,} benign; "
        f"scores {v['source']}). The regular and challenge test sets ({s['test_scores']}) are used only afterwards.",
        "A chunk is flagged when its score is at least τ. Recall = TP / (TP + FN); FPR = FP / (FP + TN); constraint: "
        f"validation FPR ≤ {100 * s['fpr_limit']:g}%.",
        f"Thresholds: the requested values {', '.join(tfmt(t) for t in req)}, and a sweep over the whole score range from "
        f"{axis_end(x_lo)} to {axis_end(x_hi)} in steps of 0.1 on the log-odds scale ln(τ / (1 − τ)). The detector's scores "
        "lie very close to 0 and 1, so the curves can only be seen on a log-odds axis.",
        "Validation-optimal interval: all thresholds with the highest validation recall under the FPR constraint and, among "
        "those, the lowest FPR. No validation score lies inside it, so every threshold in it makes the same decision on every "
        "validation chunk. Its midpoint is the threshold farthest from the nearest benign and nearest injected score.",
        "Margins: the distance from a threshold to the lower (benign) and upper (injected) end of the interval, in decades of "
        "odds, i.e. log10 of the odds ratio. A margin of 2 means a factor of 100 in the odds τ / (1 − τ).",
        f"Rule for the requested thresholds: {rule}.",
    ])

    rep.h2("Validation results (used for the selection)")
    rows = []
    for r in s["grid"]:
        note = ", ".join(x for x, flag in (("τ*", r["selected"]), ("current", r["current"])) if flag)
        rows.append([tfmt(r["threshold"]), f"{r['TP']:,}", f"{r['FN']:,}", f"{r['FP']:,}", f"{r['TN']:,}", pv(r["recall"], 2),
                     pv(r["fpr"], 2), "yes" if r["feasible"] else "no", margin(r["lower_margin"]),
                     margin(r["upper_margin"]), note])
    rep.table("Validation results at the requested thresholds, with the margins to both ends of the validation-optimal "
              "interval", ["τ", "TP", "FN", "FP", "TN", "Recall (%)", "FPR (%)", "FPR ≤ 2%", "Lower margin", "Upper margin",
                           "Note"], rows, [1.4, 1.3, 1.1, 1.1, 1.3, 1.7, 1.5, 1.4, 1.8, 1.8, 1.7],
              note="Margins in decades of odds to the lower end (benign side) and the upper end (injected side) of the "
                   "validation-optimal interval; a negative margin means the threshold lies outside the interval on that side.")
    n_tab_grid = rep.tables
    n_fig_val = None
    fig, source = pick_figure(out, "fig_threshold_sensitivity")
    if fig:
        rep.figure(fig, f"Validation recall (orange) and FPR (blue) for every threshold from {axis_end(x_lo)} to "
                        f"{axis_end(x_hi)} on a log-odds axis (top), and the validation scores of benign (blue) and injected "
                        "(orange) chunks on the same axis (bottom, log counts). Green: validation-optimal interval; grey: "
                        f"requested thresholds {tfmt(min(req))}–{tfmt(max(req))} (dots); dashed: τ = {cur:g}; dotted: midpoint "
                        f"of the interval ({source} figure).", 11.0)
        n_fig_val = rep.figures
    rows = [[tfmt(r["threshold"]) + (" (current)" if r["current"] else ""), f"{r['TP']:,}", f"{r['FN']:,}", f"{r['FP']:,}",
             f"{r['TN']:,}", pv(r["recall"], 2), pv(r["fpr"], 2), "yes" if r["feasible"] else "no"] for r in s["range"]]
    rep.table("Validation results across the whole score range", ["τ", "TP", "FN", "FP", "TN", "Recall (%)", "FPR (%)",
                                                                   "FPR ≤ 2%"], rows, [2.8, 1.6, 1.6, 1.6, 1.6, 2.2, 2.0, 1.8])
    n_tab_range = rep.tables
    rep.h2("Interpretation")
    for text in interpret_validation(s, n_tab_grid, n_tab_range, n_fig_val):
        rep.para(text)

    rep.h2("Final evaluation on the test sets (not used for the selection)")
    fe = s["final_evaluation"]
    by = {(r["split"], r["threshold"]): r for r in fe}
    rows = []
    for t in sorted({r["threshold"] for r in fe}):
        a, b = by[("regular", t)], by[("challenge", t)]
        rows.append([tfmt(t) if "requested" in a["role"] or "current" in a["role"] else sfmt(t), roles(a["role"]),
                     f"{a['FN']:,} / {a['FP']:,}", pv(a["recall"], 2), f"{b['FN']:,} / {b['FP']:,}",
                     f"{pv(b['recall'], 2)} {ci(b['recall_lo'], b['recall_hi'], 2)}", pv(b["fpr"], 2)])
    rep.table("Test-set results at the same thresholds (evaluation only)",
              ["τ", "Role", "Regular FN / FP", "Regular recall (%)", "Challenge FN / FP", "Challenge recall (%) [95% CI]",
               "Challenge FPR (%)"], rows, [1.6, 2.4, 2.2, 2.0, 2.2, 3.8, 2.0])
    fig, source = pick_figure(out, "fig_threshold_evaluation")
    if fig:
        rep.figure(fig, "Recall (top, zoomed) and FPR (bottom) on the regular test (blue) and the challenge set (orange) over "
                        f"the same threshold range. Evaluation only: these sets were not used to choose τ. Shading and lines "
                        f"as in Figure {n_fig_val} ({source} figure).", 11.0)
    fam = s.get("family_recall") or []
    if fam:
        ts = sorted({r["threshold"] for r in fam})
        names = sorted({r["family"] for r in fam})
        byf = {(r["family"], r["threshold"]): r for r in fam}
        heads = [tfmt(t) if "requested" in byf[(names[0], t)]["role"] or "current" in byf[(names[0], t)]["role"]
                 else f"{sfmt(t)} (mid.)" for t in ts]
        rows = [[n.replace("_", " "), f"{byf[(n, ts[0])]['n']:,}", *[f"{byf[(n, t)]['missed']:,}" for t in ts]] for n in names]
        rows.append(["All withheld families", f"{sum(byf[(n, ts[0])]['n'] for n in names):,}",
                     *[f"{sum(byf[(n, t)]['missed'] for n in names):,}" for t in ts]])
        rep.table("Missed challenge injections per withheld attack family at each threshold (evaluation only)",
                  ["Family", "n", *heads], rows, [3.4, 1.2, *[round(11.6 / len(ts), 2)] * len(ts)])
    rep.h2("Interpretation")
    for text in interpret_evaluation(s):
        rep.para(text)

    rep.h2("How the threshold can be described in the paper")
    for text in suggested_wording(s, n_tab_grid, n_tab_range):
        rep.para(text)


def interpret_validation(s, n_tab_grid, n_tab_range, n_fig):
    v, band, cur, c = s["validation"], s["optimal_interval"], s["current_threshold"], s["current"]
    req = [r for r in s["grid"] if r["requested"]]
    first, low_t, high_t = req[0], tfmt(req[0]["threshold"]), tfmt(req[-1]["threshold"])
    fig = f"Figure {n_fig}" if n_fig else "the full-range sweep"
    out = []
    if all(r["TP"] == first["TP"] and r["FP"] == first["FP"] for r in req):
        out.append(f"Between {low_t} and {high_t} the validation results do not change (Table {n_tab_grid}): every requested "
                   f"threshold flags {first['TP']:,} of {first['positives']:,} injected and {first['FP']:,} of "
                   f"{first['negatives']:,} benign chunks (recall {p(first['recall'], 2)}, FPR {p(first['fpr'], 2)}). This is "
                   f"not a plotting error. No validation chunk has a score between {low_t} and {high_t}, so moving the threshold "
                   "within this range cannot change any decision; a flat line is the correct result for this range.")
    else:
        rec, fpr = [r["recall"] for r in req], [r["fpr"] for r in req]
        out.append(f"Across the requested thresholds, validation recall ranges from {p(min(rec), 2)} to {p(max(rec), 2)} and "
                   f"the FPR from {p(min(fpr), 2)} to {p(max(fpr), 2)} (Table {n_tab_grid}); "
                   f"{sum(r['feasible'] for r in req)} of {len(req)} thresholds satisfy FPR ≤ {100 * s['fpr_limit']:g}%.")
    rng = s["range"]
    if band and band["separable"]:
        text = (f"Recall and FPR change only outside this range ({fig}). The highest score of a benign validation chunk "
                f"is {sfmt(v['benign_max'])} and the lowest score of an injected one {sfmt(v['injected_min'])}: the FPR rises "
                "only for thresholds below the first value, and recall falls only for thresholds above the second.")
    else:
        text = (f"Over the whole score range ({fig}) the two classes overlap: benign scores reach {sfmt(v['benign_max'])} and "
                f"injected scores go down to {sfmt(v['injected_min'])}.")
    text += (f" At τ = {tfmt(rng[0]['threshold'])}, {p(rng[0]['fpr'], 1)} of the benign chunks are flagged; at τ = "
             f"{tfmt(rng[-1]['threshold'])}, recall is {p(rng[-1]['recall'], 1)} (Table {n_tab_range}).")
    if band and band.get("fpr_edge") is not None:
        text += f" The {100 * s['fpr_limit']:g}% FPR limit is exceeded only for τ ≤ {sfmt(band['fpr_edge'])}."
    out.append(text)
    if band:
        text = (f"The validation-optimal interval is ({sfmt(band['lo'])}, {sfmt(band['hi'])}]: every threshold in it gives "
                f"recall {p(band['recall'], 2)} and FPR {p(band['fpr'], 2)} on validation. It is "
                f"{band['width_decades']:.1f} decades of odds wide, and its midpoint is {sfmt(band['midpoint'])}.")
        if c["in_optimal_interval"]:
            text += (f" The current threshold {cur:g} lies inside the interval, {c['lower_margin']:.1f} decades of odds above "
                     f"its lower end and {c['upper_margin']:.1f} below its upper end. On validation it is therefore an optimal "
                     "threshold, but not the only one: every threshold in the interval performs identically. The margins "
                     f"measure its robustness: a benign chunk would need {c['lower_margin']:.1f} decades higher odds than the "
                     f"most suspicious benign validation chunk to be flagged, and an injected chunk {c['upper_margin']:.1f} "
                     "decades lower odds than the least suspicious injected one to be missed.")
            if abs(c["upper_margin"] - c["lower_margin"]) >= 0.5:
                side = "injected" if c["upper_margin"] > c["lower_margin"] else "benign"
                text += f" The threshold therefore leaves more room on the {side} side."
        else:
            text += (f" The current threshold {cur:g} lies outside the interval: on validation it gives recall "
                     f"{p(c['recall'], 2)} and FPR {p(c['fpr'], 2)}.")
        out.append(text)
    tau, tied = s["selected_threshold"], s["tied_thresholds"]
    if tau is None:
        out.append(f"None of the requested thresholds meets FPR ≤ {100 * s['fpr_limit']:g}% on validation.")
    elif len(tied) > 1:
        text = (f"{len(tied)} of the {len(req)} requested thresholds tie on validation; the rule picks the one closest to the "
                f"midpoint of the interval, τ* = {tfmt(tau)}.")
        if abs(tau - cur) < 1e-9:
            text += f" This is the current threshold {cur:g}."
        elif any(abs(t - cur) < 1e-9 for t in tied):
            text += (f" The current threshold {cur:g} gives identical validation results, so the validation data do not prefer "
                     "either value: the difference between them is a tie-break, not a difference in performance.")
        out.append(text)
    else:
        r_tau = next(r for r in s["grid"] if r["threshold"] == tau)
        r_cur = next(r for r in s["grid"] if r["current"])
        out.append(f"The rule selects τ* = {tfmt(tau)}, with validation recall {p(r_tau['recall'], 2)} and FPR "
                   f"{p(r_tau['fpr'], 2)}" + ("." if abs(tau - cur) < 1e-9 else
                                             f" (current threshold {cur:g}: {p(r_cur['recall'], 2)} and {p(r_cur['fpr'], 2)})."))
    return out


def interpret_evaluation(s):
    fe, cur = s["final_evaluation"], s["current_threshold"]
    out = []
    for split in ("regular", "challenge"):
        rows = [r for r in fe if r["split"] == split and "requested" in r["role"]]
        a, b = rows[0], rows[-1]
        if all(r["FN"] == a["FN"] and r["FP"] == a["FP"] for r in rows):
            text = (f"{SPLIT_LABELS[split]}: every requested threshold gives {a['FN']:,} missed injections and {a['FP']:,} false "
                    f"alarms (recall {p(a['recall'], 2)}, FPR {p(a['fpr'], 2)}).")
        else:
            text = (f"{SPLIT_LABELS[split]}: here the requested thresholds are not equivalent. Missed injections go from "
                    f"{a['FN']:,} at τ = {tfmt(a['threshold'])} to {b['FN']:,} at τ = {tfmt(b['threshold'])}, and false alarms "
                    f"from {a['FP']:,} to {b['FP']:,}.")
            r_cur = next((r for r in fe if r["split"] == split and "current" in r["role"]), None)
            if r_cur:
                text += (f" At {cur:g}: {r_cur['FN']:,} missed and {r_cur['FP']:,} false alarms (recall "
                         f"{p(r_cur['recall'], 2)}, FPR {p(r_cur['fpr'], 2)}).")
        r_mid = next((r for r in fe if r["split"] == split and "midpoint" in r["role"] and "current" not in r["role"]), None)
        if r_mid:
            text += f" At the midpoint of the validation interval ({sfmt(r_mid['threshold'])}): {r_mid['FN']:,} missed and " \
                    f"{r_mid['FP']:,} false alarms."
        out.append(text)
    out.append("These test-set numbers were examined only after the selection and were not used to choose the threshold. The "
               "challenge set contains attack families that are absent from training and validation, so it shows how the "
               "threshold behaves under a shift that the validation set cannot reveal. Moving τ because of these numbers would "
               "tune it on test data; a threshold chosen for unseen families needs a validation set that contains such families.")
    return out


def suggested_wording(s, n_tab_grid, n_tab_range):
    v, band, cur, c = s["validation"], s["optimal_interval"], s["current_threshold"], s["current"]
    if not band:
        return ["No threshold meets the FPR constraint on validation, so the threshold cannot be justified from validation."]
    if not c["in_optimal_interval"]:
        out = [f"{cur:g} lies outside the validation-optimal interval for this checkpoint, so it cannot be described as the "
               f"validation-selected threshold. Report the selected value (τ* = {tfmt(s['selected_threshold'])} among the "
               f"requested thresholds; midpoint of the interval {sfmt(band['midpoint'])}) or state why {cur:g} is kept."]
    else:
        low_ref = "the highest benign score" if band["separable"] else "the lower end of this interval"
        high_ref = "the lowest injected score" if band["separable"] else "its upper end"
        text = ("Suggested text: \"The detection threshold was chosen on the validation set under the constraint FPR ≤ "
                f"{100 * s['fpr_limit']:g}%. ")
        if band["separable"]:
            text += (f"On validation the detector separates the two classes completely: every benign chunk scores at most "
                     f"{sfmt(v['benign_max'])} and every injected chunk at least {sfmt(v['injected_min'])}, so every threshold "
                     f"in ({sfmt(band['lo'])}, {sfmt(band['hi'])}] attains {p(band['recall'], 0)} recall at "
                     f"{p(band['fpr'], 0)} FPR (Fig. X). ")
        else:
            text += (f"Every threshold in ({sfmt(band['lo'])}, {sfmt(band['hi'])}] attains the highest validation recall "
                     f"({p(band['recall'], 2)}) at the lowest admissible FPR ({p(band['fpr'], 2)}) (Fig. X). ")
        text += (f"We use τ = {cur:g}, which lies inside this interval, {c['lower_margin']:.1f} decades of odds above {low_ref} "
                 f"and {c['upper_margin']:.1f} decades below {high_ref}. Because all thresholds in the interval perform "
                 "identically on validation, the validation data confirm this choice but do not single it out")
        if c["upper_margin"] > c["lower_margin"]:
            text += ("; we keep τ on the benign side of the interval because, in the hybrid design, a false positive only sends "
                     "a benign chunk through the sanitizer, whereas a false negative lets an injection through unchanged")
        text += ".\""
        out = [text, "Replace Fig. X with the number of the threshold figure in the paper. The numbers in the text come from "
                     f"Tables {n_tab_grid} and {n_tab_range} and the interpretation above."]
    if s.get("checkpoint_threshold_differs"):
        out.append("Use this text only after repeating the analysis with the checkpoint that selected the threshold (see the "
                   "note at the start of this section); the numbers above belong to the checkpoint analysed here.")
    return out


# --------------------------------------------------------------------------- section 3: distributions
def distribution_section(rep: Report, out: Path):
    rep.h1("3  Score distributions of benign and injected chunks")
    f = out / "score_distribution_summary.json"
    if not f.exists():
        rep.para("Not run yet. Run: .venv\\Scripts\\python.exe supplementary\\score_distributions.py", italic=True)
        return
    s = json.loads(f.read_text())
    rep.para(f"Scores are the saved per-chunk predictions of the trained detector on the regular test and the challenge "
             f"set. Benign and injected chunks are shown as separate distributions; the dashed line marks τ = {s['threshold']:g}.")
    fig, source = pick_figure(out, "fig_score_distributions")
    if fig:
        rep.figure(fig, f"Detector scores of benign (top, blue) and injected chunks (bottom, orange) on the regular test (left) "
                        f"and challenge set (right); log-scaled counts, bins of {s['bin_width']:g}; dashed line: τ = "
                        f"{s['threshold']:g} ({source} figure).", 11.5)
    rows = []
    for split in ("regular", "challenge"):
        info = s["splits"][split]
        for cls in ("benign", "injected"):
            st = info[cls]
            side = (f"{st['at_or_above_threshold']:,} ≥ τ" if cls == "benign" else f"{st['below_threshold']:,} < τ")
            rows.append([SPLIT_LABELS[split], cls, f"{st['n']:,}", f"{st['median']:.4f}", f"{st['p01']:.4f}",
                         f"{st['p99']:.4f}", f"{st['min']:.4f}", f"{st['max']:.4f}", side])
    rep.table("Score summary per test set and class", ["Test set", "Class", "n", "Median", "1st pct.", "99th pct.", "Min",
                                                       "Max", "Wrong side of τ"], rows,
              [2.0, 1.6, 1.4, 1.7, 1.7, 1.7, 1.7, 1.7, 2.9])
    rep.h2("Interpretation")
    for text in interpret_distributions(s):
        rep.para(text)


def interpret_distributions(s):
    out, tau = [], s["threshold"]
    for split in ("regular", "challenge"):
        info = s["splits"][split]
        b, i = info["benign"], info["injected"]
        text = f"{SPLIT_LABELS[split]}: "
        if b["max"] < tau <= i["min"]:
            text += (f"every benign chunk scored at most {b['max']:.4f} and every injected chunk at least {i['min']:.4f}, so "
                     f"the two classes are completely separated and τ = {tau:g} lies inside the gap.")
        else:
            text += (f"{info['false_alarms']:,} benign chunks reach τ and {info['misses']:,} injected chunks fall below it; "
                     f"the highest benign score is {b['max']:.4f} and the lowest injected score {i['min']:.4f}.")
        text += (f" The median benign score is {b['median']:.4f} and the median injected score {i['median']:.4f} "
                 f"(ROC AUC {info['roc_auc']:.4f}).")
        if info["misses"]:
            k, b = info["misses_between_0.10_and_threshold"], info["benign_between_0.10_and_threshold"]
            text += (f" The {info['misses']:,} missed injections have a median score of {info['median_score_of_misses']:.4f} "
                     f"(highest {info['max_score_of_misses']:.4f}); {k:,} of them lie between 0.10 and τ, so on this set a "
                     f"threshold of 0.10 would catch {k:,} of the {info['misses']:,} misses and flag {b:,} more benign chunks "
                     "(Section 2 gives the counts for every threshold).")
        out.append(text)
    out.append("This description uses the test sets; the threshold itself is chosen on the validation set (Section 2).")
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--project-root")
    ap.add_argument("--out-dir")
    ap.add_argument("--output", help="default: results/supplementary/Supplementary_Experiments_Report.docx")
    args = ap.parse_args(argv)
    try:
        import docx  # noqa: F401
    except ImportError:
        raise SystemExit("python-docx is missing. Install it once with: .venv\\Scripts\\pip.exe install python-docx")
    root = C.find_project_root(args.project_root)
    out = C.out_dir(root, args.out_dir)
    rep = Report("Supplementary Experiments: Defense Ablation, Threshold Sensitivity and Score Separation")
    rep.para(f"Layer 1 DistilBERT detector and Layer 2 clause sanitizer. Generated on {time.strftime('%d %B %Y')} from "
             f"the files in results/supplementary. Data: the 10,000-chunk validation, regular-test and challenge "
             f"partitions of the 100,000-chunk benchmark.", italic=True, size=10)
    ablation_section(rep, out)
    threshold_section(rep, out)
    distribution_section(rep, out)
    path = Path(args.output) if args.output else out / "Supplementary_Experiments_Report.docx"
    rep.save(path)
    print(f"Wrote {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
